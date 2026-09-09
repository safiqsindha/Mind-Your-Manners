"""Verifies the network-isolation claim in harness/study2/sandbox.py's
docstring: model-generated code executed via execute_python_on_workbook()
cannot reach the network when sandbox_isolation_mode() == "namespace".
Skipped (not silently passed) when `unshare` isn't on the host, so a CI
runner without it can't produce a false "isolation works" signal.
"""
from pathlib import Path

import openpyxl
import pytest

from harness.study2.sandbox import execute_python_on_workbook, sandbox_isolation_mode

pytestmark = pytest.mark.skipif(
    sandbox_isolation_mode() != "namespace",
    reason="`unshare` not available on this host -- network isolation cannot be verified here",
)


def _make_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(path)


def test_network_is_actually_unreachable(tmp_path: Path):
    in_path = tmp_path / "in.xlsx"
    _make_workbook(in_path)

    code = (
        "import socket\n"
        "try:\n"
        "    socket.create_connection(('1.1.1.1', 80), timeout=3)\n"
        "    print('REACHABLE')\n"
        "except OSError as e:\n"
        "    print('BLOCKED', e)\n"
    )
    result = execute_python_on_workbook(code, in_path, tmp_path / "work")
    assert result.network_isolated is True
    assert "BLOCKED" in result.stdout
    assert "REACHABLE" not in result.stdout


def test_isolated_execution_still_reads_and_writes_the_workbook(tmp_path: Path):
    in_path = tmp_path / "in.xlsx"
    _make_workbook(in_path)

    code = (
        "import openpyxl\n"
        "wb = openpyxl.load_workbook(WORKBOOK_PATH)\n"
        "wb.active['B1'] = wb.active['A1'].value * 2\n"
        "wb.save(OUTPUT_PATH)\n"
    )
    result = execute_python_on_workbook(code, in_path, tmp_path / "work2")
    assert result.network_isolated is True
    assert result.returncode == 0
    assert result.output_workbook_path is not None
    out_wb = openpyxl.load_workbook(result.output_workbook_path)
    assert out_wb.active["B1"].value == 2
