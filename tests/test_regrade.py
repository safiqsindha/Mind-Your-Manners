"""Tests for re-grading a finished run from its raw per-call log.

The value of regrade rests entirely on one property: replaying stored
responses must reconstruct the SAME answer the original run graded. If the
replay drifted from run_react_multi_round's bookkeeping -- which turn's
output is final, when a trajectory stops, what counts as hitting the turn
limit -- a re-grade would silently score something the models never
produced, which is worse than having no regrade at all.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from harness.study2.regrade import (
    TrajectoryKey,
    group_trajectories,
    load_raw_rows,
    rebuild_trajectory,
)

SAVE_42 = "```python\nimport openpyxl\nwb = openpyxl.load_workbook(WORKBOOK_PATH)\nwb.active['B1'] = 42\nwb.save(OUTPUT_PATH)\n```"
SAVE_99 = "```python\nimport openpyxl\nwb = openpyxl.load_workbook(WORKBOOK_PATH)\nwb.active['B1'] = 99\nwb.save(OUTPUT_PATH)\n```"
NO_CODE = "Let me think about this for a moment."
CRASHES = "```python\nraise RuntimeError('boom')\n```"


def _row(turn: int, text: str, *, task="t1", model="m", tone="none", trial=0, refused=False):
    return {
        "model_key": model, "item_id": task, "tone_level": tone, "trial": trial,
        "response_text": text, "refused": refused, "extra": {"turn": turn},
    }


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    p = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(p)
    return p


def _key(task="t1"):
    return TrajectoryKey(model_key="m", task_id=task, tone_level="none", trial=0)


# --- Grouping --------------------------------------------------------------

def test_rows_group_into_trajectories_by_model_task_tone_trial():
    rows = [
        _row(0, SAVE_42, task="a", tone="L1_sycophantic"),
        _row(0, SAVE_42, task="a", tone="L5_rude"),
        _row(0, SAVE_42, task="b", tone="L1_sycophantic"),
        _row(0, SAVE_42, task="a", tone="L1_sycophantic", trial=1),
    ]
    assert len(group_trajectories(rows)) == 4, "each condition is its own trajectory"


def test_turns_are_ordered_even_when_the_log_is_not():
    rows = [_row(2, SAVE_99), _row(0, SAVE_42), _row(1, NO_CODE)]
    turns = next(iter(group_trajectories(rows).values()))
    assert [(r["extra"] or {})["turn"] for r in turns] == [0, 1, 2]


def test_a_truncated_final_line_does_not_lose_the_rest(tmp_path: Path):
    """A process killed mid-write leaves a partial line; refusing to read the
    other 99.9% of the log would defeat the point of having it."""
    p = tmp_path / "raw.jsonl"
    p.write_text(
        json.dumps(_row(0, SAVE_42)) + "\n"
        + json.dumps(_row(1, SAVE_99)) + "\n"
        + '{"model_key": "m", "response_te'
    )
    assert len(load_raw_rows(p)) == 2


# --- Reconstruction fidelity ----------------------------------------------

def test_the_last_turn_that_wrote_output_is_the_one_graded(tmp_path: Path, workbook: Path):
    """Mirrors the agent loop: final_output_path is overwritten by whichever
    turn produced output most recently, so a later write wins."""
    turns = [_row(0, SAVE_42), _row(1, SAVE_99)]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd")

    assert traj.final_output_path is not None
    assert openpyxl.load_workbook(traj.final_output_path).active["B1"].value == 99


def test_a_later_crash_does_not_discard_an_earlier_good_output(tmp_path: Path, workbook: Path):
    turns = [_row(0, SAVE_42), _row(1, CRASHES)]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd")

    assert traj.final_output_path is not None
    assert openpyxl.load_workbook(traj.final_output_path).active["B1"].value == 42


def test_final_ends_the_trajectory_and_later_rows_are_ignored(tmp_path: Path, workbook: Path):
    """FINAL breaks the loop in the original, so anything logged after it was
    never part of that trajectory."""
    turns = [_row(0, SAVE_42), _row(1, "FINAL: done"), _row(2, SAVE_99)]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd")

    assert len(traj.steps) == 2
    assert traj.steps[-1].is_final
    assert openpyxl.load_workbook(traj.final_output_path).active["B1"].value == 42


def test_a_refusal_ends_the_trajectory_and_is_recorded(tmp_path: Path, workbook: Path):
    turns = [_row(0, "I won't do that", refused=True), _row(1, SAVE_42)]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd")

    assert traj.refused is True
    assert len(traj.steps) == 1
    assert traj.final_output_path is None


def test_turns_without_code_are_kept_as_steps_not_dropped(tmp_path: Path, workbook: Path):
    """The original records a step for a protocol violation and keeps going;
    dropping it here would misreport n_turns."""
    turns = [_row(0, NO_CODE), _row(1, SAVE_42)]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd")

    assert len(traj.steps) == 2
    assert traj.steps[0].code is None
    assert traj.final_output_path is not None


# --- Turn-limit bookkeeping ------------------------------------------------

def test_running_out_of_turns_without_final_sets_hit_turn_limit(tmp_path: Path, workbook: Path):
    turns = [_row(i, SAVE_42) for i in range(10)]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd", max_turns=10)
    assert traj.hit_turn_limit is True


def test_finishing_early_with_final_does_not_set_hit_turn_limit(tmp_path: Path, workbook: Path):
    turns = [_row(0, SAVE_42), _row(1, "FINAL: done")]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd", max_turns=10)
    assert traj.hit_turn_limit is False


def test_a_short_trajectory_is_not_mislabelled_when_max_turns_is_unknown(tmp_path: Path, workbook: Path):
    """Without max_turns the recorded turn count is the only evidence, so a
    trajectory that simply stopped early must not be called turn-limited."""
    turns = [_row(0, SAVE_42), _row(1, "FINAL: done")]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd", max_turns=None)
    assert traj.hit_turn_limit is False


# --- No model calls --------------------------------------------------------

def test_regrade_never_calls_a_provider(tmp_path: Path, workbook: Path, monkeypatch):
    """The entire premise is that this is free. A provider call here would
    mean re-paying for work already paid for."""
    import harness.providers.registry as registry

    def explode(*a, **kw):
        raise AssertionError("regrade must not call any provider")

    monkeypatch.setattr(registry, "get_provider", explode)
    turns = [_row(0, SAVE_42), _row(1, "FINAL: done")]
    traj = rebuild_trajectory(_key(), turns, workbook, tmp_path / "wd")
    assert traj.final_output_path is not None
