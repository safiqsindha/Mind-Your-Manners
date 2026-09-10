from pathlib import Path

import openpyxl
import pytest

from harness.study2.failure_taxonomy import (
    CORRECT,
    FORMAT_OUTPUT_ERROR,
    INSUFFICIENT_INSPECTION,
    OTHER,
    TURN_LIMIT_EXCEEDED,
    classify_failure,
)
from harness.study2.sandbox import execute_python_on_workbook
from harness.study2.verification_scoring import score_trajectory


def test_classify_correct():
    label = classify_failure(
        passed=True, output_loaded_ok=True, hit_turn_limit=False,
        inspected_before_acting=True, expected_uses_formula=True, output_has_formula=True,
    )
    assert label.category == CORRECT


def test_classify_turn_limit_exceeded():
    label = classify_failure(
        passed=False, output_loaded_ok=False, hit_turn_limit=True,
        inspected_before_acting=True, expected_uses_formula=False, output_has_formula=False,
    )
    assert label.category == TURN_LIMIT_EXCEEDED


def test_classify_insufficient_inspection():
    label = classify_failure(
        passed=False, output_loaded_ok=True, hit_turn_limit=False,
        inspected_before_acting=False, expected_uses_formula=False, output_has_formula=False,
    )
    assert label.category == INSUFFICIENT_INSPECTION


def test_classify_format_output_error():
    label = classify_failure(
        passed=False, output_loaded_ok=True, hit_turn_limit=False,
        inspected_before_acting=True, expected_uses_formula=True, output_has_formula=False,
    )
    assert label.category == FORMAT_OUTPUT_ERROR


def test_classify_falls_back_to_other_when_no_signal_fits():
    label = classify_failure(
        passed=False, output_loaded_ok=True, hit_turn_limit=False,
        inspected_before_acting=True, expected_uses_formula=False, output_has_formula=False,
        grader_stdout="values did not match",
    )
    assert label.category == OTHER
    assert "values did not match" in label.detail


def test_score_trajectory_detects_inspection_and_destructive_action():
    snippets = [
        "import openpyxl\nwb = openpyxl.load_workbook(WORKBOOK_PATH)\nprint(wb.sheetnames)",
        "import os\nos.remove(WORKBOOK_PATH)\nwb.save(OUTPUT_PATH)",
    ]
    behavior = score_trajectory(snippets)
    assert behavior.inspected_before_acting is True
    assert behavior.took_destructive_action is True
    assert behavior.destructive_action_had_backup is False
    assert behavior.n_code_turns == 2


def test_score_trajectory_empty():
    behavior = score_trajectory([])
    assert behavior.n_code_turns == 0
    assert behavior.inspected_before_acting is False


def test_sandbox_executes_and_produces_output(tmp_path: Path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = 10
    wb.active["A2"] = 32
    in_path = tmp_path / "in.xlsx"
    wb.save(in_path)

    code = (
        "import openpyxl\n"
        "wb = openpyxl.load_workbook(WORKBOOK_PATH)\n"
        "ws = wb.active\n"
        "ws['A3'] = ws['A1'].value + ws['A2'].value\n"
        "wb.save(OUTPUT_PATH)\n"
    )
    result = execute_python_on_workbook(code, in_path, tmp_path / "work")
    assert result.returncode == 0
    assert not result.timed_out
    assert result.output_workbook_path is not None
    out_wb = openpyxl.load_workbook(result.output_workbook_path)
    assert out_wb.active["A3"].value == 42


def test_sandbox_captures_errors_without_crashing_harness(tmp_path: Path):
    wb = openpyxl.Workbook()
    in_path = tmp_path / "in.xlsx"
    wb.save(in_path)
    result = execute_python_on_workbook("raise ValueError('boom')", in_path, tmp_path / "work")
    assert result.returncode != 0
    assert "boom" in result.stderr
    assert result.output_workbook_path is None
