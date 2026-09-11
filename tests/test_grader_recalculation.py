"""Regression test for the bug found via gating check (see grader.py module
docstring): a ground-truth answer file with an uncached formula reads as
None and fails comparison against an objectively-correct model output,
unless the ground truth is ALSO recalculated -- not just the model's
output. Uses a minimal fake SpreadsheetBench checkout (just enough of
evaluation/evaluation.py to exercise real cell comparison) rather than a
full clone, but exercises the real LibreOffice recalculation path -- skipped
if LibreOffice isn't available."""
import shutil
import textwrap
from pathlib import Path

import openpyxl
import pytest

from harness.study2.grader import SpreadsheetBenchGrader

pytestmark = pytest.mark.skipif(
    shutil.which("soffice") is None and shutil.which("libreoffice") is None,
    reason="LibreOffice not available on this host",
)

_FAKE_EVALUATION_PY = textwrap.dedent(
    """
    import openpyxl

    def compare_workbooks(gt_file, proc_file, instruction_type, answer_position):
        wb_gt = openpyxl.load_workbook(gt_file, data_only=True)
        wb_proc = openpyxl.load_workbook(proc_file, data_only=True)
        gt_val = wb_gt.active[answer_position].value
        proc_val = wb_proc.active[answer_position].value
        return (gt_val == proc_val), ""
    """
)


@pytest.fixture
def fake_repo(tmp_path: Path) -> Path:
    eval_dir = tmp_path / "evaluation"
    eval_dir.mkdir()
    (eval_dir / "evaluation.py").write_text(_FAKE_EVALUATION_PY)
    return tmp_path


def test_uncached_ground_truth_formula_is_recalculated_before_comparison(fake_repo: Path, tmp_path: Path):
    # Ground truth: A1=6, A2=7, A3=formula (uncached -- as if just written by openpyxl, never opened in Excel).
    answer_wb = openpyxl.Workbook()
    answer_wb.active["A1"] = 6
    answer_wb.active["A2"] = 7
    answer_wb.active["A3"] = "=A1*A2"
    answer_path = tmp_path / "answer.xlsx"
    answer_wb.save(answer_path)

    # A perfectly correct model output: computed the right value directly.
    output_wb = openpyxl.Workbook()
    output_wb.active["A3"] = 42
    output_path = tmp_path / "output.xlsx"
    output_wb.save(output_path)

    grader = SpreadsheetBenchGrader(fake_repo)
    result = grader.evaluate_task(
        task_id="t1",
        instruction_type="Cell-Level Manipulation",
        answer_position="A3",
        output_paths=[output_path],
        answer_paths=[answer_path],
        recalculate=True,
    )
    assert result.passed, "a correct output must not fail because the ground truth's own formula was uncached"


def test_answer_file_recalculation_is_memoized(fake_repo: Path, tmp_path: Path, monkeypatch):
    answer_wb = openpyxl.Workbook()
    answer_wb.active["A1"] = 42
    answer_path = tmp_path / "answer.xlsx"
    answer_wb.save(answer_path)
    output_wb = openpyxl.Workbook()
    output_wb.active["A1"] = 42
    output_path = tmp_path / "output.xlsx"
    output_wb.save(output_path)

    grader = SpreadsheetBenchGrader(fake_repo)

    calls = []
    import harness.study2.grader as grader_module

    real_recalc = grader_module.recalculate_with_libreoffice

    def spy(paths, soffice_bin=None):
        calls.append([str(p) for p in paths])
        return real_recalc(paths, soffice_bin=soffice_bin)

    monkeypatch.setattr(grader_module, "recalculate_with_libreoffice", spy)

    before = answer_path.read_bytes()

    for _ in range(3):
        grader.evaluate_task("t1", "Cell-Level Manipulation", "A1", [output_path], [answer_path], recalculate=True)

    # The original ground-truth file must never be handed to the in-place
    # recalculator. It used to be, which silently overwrote SpreadsheetBench's
    # shipped answer with a LibreOffice round-trip -- permanently, since the
    # tarball is never re-extracted once the directory exists.
    assert all(str(answer_path) not in c for c in calls), (
        f"ground truth was passed to the in-place recalculator: {calls}"
    )
    assert answer_path.read_bytes() == before, "ground-truth answer file was modified"

    # The expensive conversion still happens only once across repeated calls
    # (the point of the old memo), now via an on-disk cache that also holds
    # across processes rather than dying with the instance.
    cache_dir = grader.recalc_cache_dir
    converted = sum(1 for c in calls if any(str(cache_dir) in p for p in c))
    assert converted == 1, f"cached copy should be converted once, not per call: {calls}"


def test_a_second_grader_instance_reuses_the_cached_conversion(tmp_path, fake_repo, monkeypatch):
    """The old memo was per-instance, so every new process re-converted files
    already converted -- a LibreOffice round-trip stacked on the previous
    one, compounding indefinitely."""
    import openpyxl

    import harness.study2.grader as grader_module

    answer_wb = openpyxl.Workbook()
    answer_wb.active["A1"] = 42
    answer_path = tmp_path / "answer.xlsx"
    answer_wb.save(answer_path)
    output_wb = openpyxl.Workbook()
    output_wb.active["A1"] = 42
    output_path = tmp_path / "output.xlsx"
    output_wb.save(output_path)

    SpreadsheetBenchGrader(fake_repo).evaluate_task(
        "t1", "Cell-Level Manipulation", "A1", [output_path], [answer_path], recalculate=True
    )

    calls = []
    real_recalc = grader_module.recalculate_with_libreoffice

    def spy(paths, soffice_bin=None):
        calls.append([str(p) for p in paths])
        return real_recalc(paths, soffice_bin=soffice_bin)

    monkeypatch.setattr(grader_module, "recalculate_with_libreoffice", spy)

    # A fresh instance, as a new process would build.
    grader2 = SpreadsheetBenchGrader(fake_repo)
    grader2.evaluate_task(
        "t1", "Cell-Level Manipulation", "A1", [output_path], [answer_path], recalculate=True
    )

    cache_dir = grader2.recalc_cache_dir
    answer_conversions = sum(1 for c in calls if any(str(cache_dir) in p for p in c))
    assert answer_conversions == 0, f"cached conversion should be reused across instances: {calls}"
