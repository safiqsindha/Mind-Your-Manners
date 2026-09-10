"""Verifies harness/study2/grader.py:recalculate_with_libreoffice() actually
recalculates formulas, rather than assuming it does. Two real bugs were
found and fixed here (see grader.py module docstring): missing LibreOffice
component packages in this build's container, and a silent in-place-write
failure that the previous version's success check didn't catch. Skipped
(not silently passed) when soffice/libreoffice isn't on PATH."""
import shutil
from pathlib import Path

import openpyxl
import pytest

from harness.study2.grader import recalculate_with_libreoffice

pytestmark = pytest.mark.skipif(
    shutil.which("soffice") is None and shutil.which("libreoffice") is None,
    reason="LibreOffice not available on this host -- cannot verify recalculation",
)


def test_recalculates_a_real_formula(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 7
    ws["A2"] = 6
    ws["A3"] = "=A1*A2"
    path = tmp_path / "t.xlsx"
    wb.save(path)

    # Before recalculation, openpyxl with data_only=True has no cached value.
    before = openpyxl.load_workbook(path, data_only=True)
    assert before.active["A3"].value is None

    ok, msg = recalculate_with_libreoffice([path])
    assert ok is True, msg

    after = openpyxl.load_workbook(path, data_only=True)
    assert after.active["A3"].value == 42


def test_reports_failure_for_a_missing_file(tmp_path: Path):
    ok, msg = recalculate_with_libreoffice([tmp_path / "does_not_exist.xlsx"])
    assert ok is False
    assert msg
