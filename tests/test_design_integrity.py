"""Tests for three design-level holes found in a pre-core-run review.

Unlike most of this suite these are not about crashes. Each one could have
produced a plausible-looking NUMBER that was wrong, which is worse: a crash
announces itself, a confound does not.
"""
from __future__ import annotations

import collections
from pathlib import Path
from dataclasses import replace
from unittest.mock import patch

import openpyxl
import pytest

from harness.providers.base import ModelConfig
from harness.study2.runner import _run_tag
from harness.study2.sandbox import execute_python_on_workbook

LIST_SIBLINGS = """
import os
d = os.path.dirname(WORKBOOK_PATH)
print("SIBLINGS:", sorted(os.listdir(d)))
"""


def _model(key: str) -> ModelConfig:
    return ModelConfig(
        key=key, provider="mock", model_id=f"x/{key}",
        display_name=key, temperature=0.0, max_tokens=512,
    )


# --- 1. Answer-key isolation ----------------------------------------------
# SpreadsheetBench stores ground truth beside the input:
# spreadsheet/59196/1_59196_input.xlsx next to 1_59196_answer.xlsx. With
# WORKBOOK_PATH pointing into the dataset, one os.listdir reached the answer,
# and the sandbox does not restrict filesystem access.

def _dataset_dir(tmp_path: Path) -> Path:
    """A directory shaped like SpreadsheetBench's: input beside answer."""
    d = tmp_path / "spreadsheet" / "59196"
    d.mkdir(parents=True)
    for name, value in (("1_59196_input.xlsx", 1), ("1_59196_answer.xlsx", 999)):
        wb = openpyxl.Workbook()
        wb.active["A1"] = value
        wb.save(d / name)
    return d


def test_model_code_cannot_see_the_answer_file(tmp_path: Path):
    """The whole point: a model listing its workbook's directory must not
    find the ground truth there."""
    data = _dataset_dir(tmp_path)

    result = execute_python_on_workbook(
        LIST_SIBLINGS, data / "1_59196_input.xlsx", tmp_path / "work"
    )

    assert result.returncode == 0, result.stderr
    assert "_answer" not in result.stdout, (
        f"answer key reachable from the sandbox: {result.stdout}"
    )


def test_the_workbook_the_model_reads_is_a_copy_not_the_dataset_file(tmp_path: Path):
    data = _dataset_dir(tmp_path)
    workdir = tmp_path / "work"

    execute_python_on_workbook(
        "print('PATH:', WORKBOOK_PATH)", data / "1_59196_input.xlsx", workdir
    )

    copied = workdir / "1_59196_input.xlsx"
    assert copied.exists(), "input should be copied into the workdir"
    assert copied.parent == workdir


def test_the_copy_has_the_same_contents_as_the_original(tmp_path: Path):
    """Isolation must not change what the model is working on."""
    data = _dataset_dir(tmp_path)
    src = data / "1_59196_input.xlsx"

    result = execute_python_on_workbook(
        "import openpyxl\nprint('VAL:', openpyxl.load_workbook(WORKBOOK_PATH).active['A1'].value)",
        src, tmp_path / "work",
    )

    assert "VAL: 1" in result.stdout, result.stdout
    assert (tmp_path / "work" / src.name).read_bytes() == src.read_bytes()


def test_the_original_dataset_file_is_not_modified(tmp_path: Path):
    """Model code writing to WORKBOOK_PATH must hit the copy, never the
    benchmark's own input file."""
    data = _dataset_dir(tmp_path)
    src = data / "1_59196_input.xlsx"
    before = src.read_bytes()

    execute_python_on_workbook(
        "import openpyxl\nwb = openpyxl.load_workbook(WORKBOOK_PATH)\n"
        "wb.active['A1'] = 'CLOBBERED'\nwb.save(WORKBOOK_PATH)",
        src, tmp_path / "work",
    )

    assert src.read_bytes() == before, "model code modified the dataset input"


def test_the_extension_is_preserved_so_xlsm_handling_still_works(tmp_path: Path):
    """Real model code branches on WORKBOOK_PATH.lower().endswith('.xlsm')."""
    src = tmp_path / "book.xlsm"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(src)

    result = execute_python_on_workbook(
        "print('EXT_OK:', WORKBOOK_PATH.lower().endswith('.xlsm'))", src, tmp_path / "work"
    )

    assert "EXT_OK: True" in result.stdout, result.stdout


# --- 2. Tone order randomisation ------------------------------------------
# A fixed L1..L7 sequence confounds tone with position-in-burst. Retry backoff
# accumulates across consecutive calls, so the last tone would systematically
# meet worse provider conditions than the first.

def _tone_sequences(tone_seed: int, n_tasks: int = 8):
    """Run the batch loop with everything stubbed, capturing tone order."""
    from harness.study2 import runner

    seen: dict[str, list[str]] = collections.defaultdict(list)

    class _Task:
        def __init__(self, tid):
            self.task_id = tid
            self.instruction_type = "Cell-Level Manipulation"
            self.instruction = "do it"
            self.answer_position = "A1"
            self.input_spreadsheet_paths = [Path("in.xlsx")]
            self.answer_spreadsheet_paths = [Path("ans.xlsx")]

    from harness.study2.agent_loop import Trajectory

    def fake_run(tracker, model, task_id, instruction, input_path, workdir, *, tone_level, trial, **kw):
        seen[task_id].append(tone_level)
        return Trajectory(task_id=task_id, tone_level=tone_level, trial=trial)

    def fake_grade(grader, task, traj, workdir):
        from harness.study2.grader import GradeResult
        return GradeResult(task.task_id, False, 3, 0, 0.0, [])

    with patch("harness.study2.runner.run_react_multi_round", side_effect=fake_run), \
         patch("harness.study2.runner._grade_trajectory", side_effect=fake_grade), \
         patch("harness.study2.runner.no_op_passes", return_value=False):
        runner.run_condition_batch(
            [_model("m")], [_Task(f"t{i}") for i in range(n_tasks)],
            grader=None, out_dir=Path("/tmp/_tone_order_probe"), phase="probe",
            budget_cap_usd=1e6, n_trials=1, tone_seed=tone_seed,
        )
    return seen


def test_tone_order_is_not_the_same_every_task():
    """The bug: identical L1..L7 for every task and trial."""
    seen = _tone_sequences(tone_seed=0)
    orders = {tuple(v) for v in seen.values()}
    assert len(orders) > 1, f"tone order identical across all tasks: {orders}"


def test_every_tone_still_runs_exactly_once_per_task():
    """Shuffling must not drop or duplicate a condition."""
    from harness.tone_wrappers import TONE_ORDER

    for task_id, tones in _tone_sequences(tone_seed=0).items():
        assert sorted(tones) == sorted(TONE_ORDER), f"{task_id} got {tones}"


def test_tone_order_is_reproducible_for_a_seed():
    assert _tone_sequences(tone_seed=7) == _tone_sequences(tone_seed=7)


def test_a_different_seed_gives_a_different_arrangement():
    assert _tone_sequences(tone_seed=1) != _tone_sequences(tone_seed=2)


# --- 3. Per-run file namespacing ------------------------------------------
# Four models run as concurrent processes shared one raw log, one records
# file, and -- worst -- one SpendTracker resume source, so each model would
# count all four models' spend against its own cap.

def _live(key: str) -> ModelConfig:
    return replace(_model(key), provider="openai_compatible")


def test_a_single_model_run_gets_its_own_tag():
    assert _run_tag([_live("gpt-luna")]) == "gpt-luna"


def test_two_single_model_runs_do_not_collide():
    """The parallel case: each process must write its own files."""
    tags = {_run_tag([_live(k)]) for k in ("gpt-luna", "glm-current", "deepseek-current", "qwen-current")}
    assert len(tags) == 4


def test_a_multi_model_run_in_one_process_shares_one_tag():
    """Correct: they are genuinely one run under one budget cap."""
    assert _run_tag([_live("a"), _live("b")]) == "multi"


# A dry run forces every model onto the mock provider but used to write to the
# same filenames as a live run. A mock `core` invocation left 48 fabricated
# rows in study2_core_gpt-luna.jsonl, which the next live run resumed its
# budget from and would have analysed alongside real results.

def test_a_dry_run_does_not_write_where_a_live_run_writes():
    assert _run_tag([_model("gpt-luna")]) != _run_tag([_live("gpt-luna")])


def test_a_dry_run_tag_says_so():
    assert _run_tag([_model("gpt-luna")]) == "gpt-luna-dryrun"


def test_one_mock_model_is_enough_to_mark_the_whole_run():
    """--dry-run mocks everything, but a partially-mocked run is still not
    live data and must not land in a live file."""
    assert _run_tag([_live("a"), _model("b")]).endswith("-dryrun")


# --- 4. The validation gate's spend log was still shared -------------------
# Fix 3 namespaced the core path's files but left the gate's SpendTracker log
# on one fixed name. The gate is the phase that runs four models AT ONCE by
# design, and SpendTracker resumes its running total from that file, so each
# model would count all four models' spend against its own $10 cap.

def _gate_tracker_path(model_key: str, out_dir: Path) -> Path:
    """Run the gate with zero tasks and report where SpendTracker was pointed."""
    from harness.study2 import runner

    seen: list[Path] = []

    class _FakeTracker:
        total_usd = 0.0

        def __init__(self, path, **kw):
            seen.append(Path(path))

        def close(self):
            pass

    with patch("harness.study2.runner.SpendTracker", _FakeTracker):
        runner.run_validation_gate(_model(model_key), [], grader=None, out_dir=out_dir)
    return seen[0]


def test_two_models_gating_do_not_share_a_spend_log(tmp_path: Path):
    a = _gate_tracker_path("gpt-luna", tmp_path)
    b = _gate_tracker_path("qwen-current", tmp_path)
    assert a != b, f"both models wrote spend to {a}"


def test_the_gate_spend_log_names_its_model(tmp_path: Path):
    path = _gate_tracker_path("deepseek-current", tmp_path)
    assert "deepseek-current" in path.name, path.name


# --- 5. Model version drift, and reports destroyed by stray runs -----------
# A model_id like "qwen/qwen3.8-flash" is a pointer the provider can repoint
# at a newer dated snapshot at any time. No response body reveals it: the
# response reports the pointer, and the served provider is unchanged. A
# mid-study roll would split the run across two models with every number
# still looking plausible.

def _cfg(key="qwen-current", model_id="qwen/qwen3.8-flash",
         canonical="qwen/qwen3.8-flash-20260826"):
    return ModelConfig(
        key=key, provider="openai_compatible", model_id=model_id,
        display_name=key, temperature=0.0, max_tokens=512,
        canonical_slug=canonical,
    )


def _catalog(*pairs):
    class _R:
        status_code = 200
        @staticmethod
        def json():
            return {"data": [{"id": i, "canonical_slug": c} for i, c in pairs]}
    return _R()


def test_matching_versions_pass_and_are_reported_back(monkeypatch):
    from harness.providers import openai_compatible as oc
    monkeypatch.setenv("OPENROUTER_API_KEY", "k")
    monkeypatch.setattr(oc.requests, "get",
                        lambda *a, **k: _catalog(("qwen/qwen3.8-flash", "qwen/qwen3.8-flash-20260826")))
    observed = oc.assert_canonical_slugs([_cfg()])
    assert observed == {"qwen-current": "qwen/qwen3.8-flash-20260826"}


def test_a_repointed_slug_halts_the_run(monkeypatch):
    """The whole point: the pointer now resolves somewhere else."""
    from harness.providers import openai_compatible as oc
    from harness.providers.base import ProviderPinViolation
    monkeypatch.setenv("OPENROUTER_API_KEY", "k")
    monkeypatch.setattr(oc.requests, "get",
                        lambda *a, **k: _catalog(("qwen/qwen3.8-flash", "qwen/qwen3.8-flash-20261102")))
    with pytest.raises(ProviderPinViolation) as e:
        oc.assert_canonical_slugs([_cfg()])
    assert "20261102" in str(e.value) and "20260826" in str(e.value)


def test_a_model_vanishing_from_the_catalog_halts_the_run(monkeypatch):
    from harness.providers import openai_compatible as oc
    from harness.providers.base import ProviderPinViolation
    monkeypatch.setenv("OPENROUTER_API_KEY", "k")
    monkeypatch.setattr(oc.requests, "get", lambda *a, **k: _catalog(("something/else", "something/else-1")))
    with pytest.raises(ProviderPinViolation):
        oc.assert_canonical_slugs([_cfg()])


def test_an_unreadable_catalog_fails_closed(monkeypatch):
    """Never start a paid run on unverified versions."""
    from harness.providers import openai_compatible as oc
    from harness.providers.base import ProviderError
    monkeypatch.setenv("OPENROUTER_API_KEY", "k")

    class _R:
        status_code = 503
    monkeypatch.setattr(oc.requests, "get", lambda *a, **k: _R())
    with pytest.raises(ProviderError):
        oc.assert_canonical_slugs([_cfg()])


def test_a_small_run_cannot_overwrite_a_large_runs_report(tmp_path: Path):
    """A stray --n-tasks 1 destroyed a completed 100-task report during
    development. Only a committed copy saved the numbers."""
    from harness.cli import _refuse_to_shrink_report
    import json as _json

    p = tmp_path / "study2_validation_gate_glm-current.json"
    p.write_text(_json.dumps({"n_tasks": 100, "n_passed": 28}))

    with pytest.raises(SystemExit):
        _refuse_to_shrink_report(p, {"n_tasks": 1, "n_passed": 0}, force=False)
    assert _json.loads(p.read_text())["n_tasks"] == 100, "the large report must survive"


def test_an_equal_or_larger_run_may_overwrite(tmp_path: Path):
    from harness.cli import _refuse_to_shrink_report
    import json as _json

    p = tmp_path / "r.json"
    p.write_text(_json.dumps({"n_tasks": 100}))
    _refuse_to_shrink_report(p, {"n_tasks": 100}, force=False)
    _refuse_to_shrink_report(p, {"n_tasks": 200}, force=False)


def test_force_overwrite_is_the_explicit_escape_hatch(tmp_path: Path):
    from harness.cli import _refuse_to_shrink_report
    import json as _json

    p = tmp_path / "r.json"
    p.write_text(_json.dumps({"n_tasks": 100}))
    _refuse_to_shrink_report(p, {"n_tasks": 1}, force=True)


# --- 6. The core phase drew from a different pool than everything else ----
# core used sample_only=False (909 tasks) while every other phase used the
# 200-task sample. Both things that make an accuracy interpretable -- the
# no-op manifest and the validation-gate baselines -- were measured on the
# sample, and neither transfers.

def _tasks(n, prefix="t"):
    class _T:
        def __init__(self, i):
            self.task_id = f"{prefix}{i}"
            self.instruction_type = "Cell-Level Manipulation"
    return [_T(i) for i in range(n)]


def test_core_draws_from_the_same_pool_as_the_gate_by_default():
    """The bug: a core draw of 50 overlapped the n=100 gate sample by 2."""
    from harness.study2.runner import select_gate_tasks

    pool = _tasks(200)
    core = {t.task_id for t in select_gate_tasks(pool, n=50, seed=0)}
    gate = {t.task_id for t in select_gate_tasks(pool, n=100, seed=0)}
    assert core <= gate, f"core sample is not inside the gate sample: {sorted(core - gate)[:5]}"


def test_a_bigger_pool_gives_a_different_draw():
    """Why the boundary mattered: same seed, same n, different pool."""
    from harness.study2.runner import select_gate_tasks

    small = {t.task_id for t in select_gate_tasks(_tasks(200), n=50, seed=0)}
    big = {t.task_id for t in select_gate_tasks(_tasks(909), n=50, seed=0)}
    assert len(small & big) < 50, "drawing from a larger pool should not reproduce the sample draw"


def test_every_phase_defaults_to_the_sample():
    """--full-dataset must be opt-in, on every staged phase."""
    import argparse
    from harness.cli import build_parser

    parser = build_parser()
    for phase in ("pilot", "core", "frontier"):
        args = parser.parse_args(["study2", phase])
        assert args.full_dataset is False, f"{phase} defaults to the full dataset"
        opted = parser.parse_args(["study2", phase, "--full-dataset"])
        assert opted.full_dataset is True


# --- 7. The primary outcome had no significance test ----------------------
# token_cost_effect_size reports a relative-variation percentage, pooled
# across tasks, with no p-value -- and cost is the study's PRIMARY outcome.
# Pooling is the defect: tasks differ enormously in how much thinking they
# demand, and that variance swamps any tone effect. Measured on a partial
# Luna core run, pooling gave p=0.81 while the same data clustered by task
# gave p=0.03.

def _token_rows(effect_on_threatening: int, n_tasks: int = 12, seed: int = 0):
    """Tasks with wildly different baseline token demand, plus a real
    threatening-tone bump. Pooled, the task spread hides the bump."""
    import random as _r
    from harness.tone_wrappers import TONE_ORDER

    rng = _r.Random(seed)
    rows = []
    for i in range(n_tasks):
        base = 200 * (i + 1) * 10          # task demand spans 2k..24k
        for tone in TONE_ORDER:
            bump = effect_on_threatening if tone == "L7_threatening" else 0
            rows.append({
                "task_id": f"t{i}", "tone_level": tone, "passed": False,
                "reasoning_tokens": base + bump + rng.randint(-50, 50),
            })
    return rows


def test_a_real_within_task_effect_is_detected():
    from harness.study2.analysis import token_cost_trend_test

    t = token_cost_trend_test(_token_rows(effect_on_threatening=400))
    assert t.p_value < 0.05, f"missed a real effect, p={t.p_value}"
    assert t.observed_slope > 0


def test_no_effect_is_not_invented():
    from harness.study2.analysis import token_cost_trend_test

    t = token_cost_trend_test(_token_rows(effect_on_threatening=0))
    assert t.p_value > 0.05, f"found an effect that is not there, p={t.p_value}"


def test_it_defaults_to_reasoning_not_total_tokens():
    """total_tokens is dominated by the prompt, whose length the tone
    wrapper changes by construction."""
    import inspect
    from harness.study2.analysis import token_cost_trend_test

    assert inspect.signature(token_cost_trend_test).parameters["value_key"].default == "reasoning_tokens"


def test_a_run_without_the_field_says_so_rather_than_scoring_zero():
    """No thinking measurement is not a measurement of no thinking."""
    from harness.study2.analysis import token_cost_trend_test

    rows = [{"task_id": "t0", "tone_level": t, "passed": False} for t in
            ("L1_sycophantic", "L4_neutral", "L7_threatening")]
    with pytest.raises(ValueError, match="reasoning_tokens"):
        token_cost_trend_test(rows)


def test_the_report_carries_both_token_trends():
    from harness.cli import _token_cost_trends

    out = _token_cost_trends(_token_rows(effect_on_threatening=400))
    assert set(out) == {"reasoning_tokens", "total_tokens"}
    assert out["reasoning_tokens"]["p_value"] < 0.05
    assert "unavailable" in out["total_tokens"], "total_tokens absent here, should say so"


def test_reasoning_tokens_are_recoverable_from_the_raw_log(tmp_path: Path):
    """The field was added mid-run. A run recorded before it must not become
    unanalysable on the study's primary outcome."""
    import json as _json
    from harness.study2.analysis import backfill_reasoning_tokens

    raw = tmp_path / "raw.jsonl"
    raw.write_text("\n".join(_json.dumps(r) for r in [
        {"item_id": "t0", "tone_level": "L4_neutral", "trial": 0, "reasoning_tokens": 100},
        {"item_id": "t0", "tone_level": "L4_neutral", "trial": 0, "reasoning_tokens": 50},
        {"item_id": "t0", "tone_level": "L7_threatening", "trial": 0, "reasoning_tokens": 900},
    ]))
    recs = [
        {"task_id": "t0", "tone_level": "L4_neutral", "trial": 0},
        {"task_id": "t0", "tone_level": "L7_threatening", "trial": 0},
    ]
    assert backfill_reasoning_tokens(recs, raw) == 2
    assert recs[0]["reasoning_tokens"] == 150, "per-call values must be summed per trajectory"
    assert recs[1]["reasoning_tokens"] == 900


def test_backfill_leaves_existing_values_alone(tmp_path: Path):
    import json as _json
    from harness.study2.analysis import backfill_reasoning_tokens

    raw = tmp_path / "raw.jsonl"
    raw.write_text(_json.dumps({"item_id": "t0", "tone_level": "L4_neutral", "trial": 0, "reasoning_tokens": 999}))
    recs = [{"task_id": "t0", "tone_level": "L4_neutral", "trial": 0, "reasoning_tokens": 42}]
    assert backfill_reasoning_tokens(recs, raw) == 0
    assert recs[0]["reasoning_tokens"] == 42


def test_backfill_survives_a_torn_final_line(tmp_path: Path):
    """A run still in flight can leave a partial line."""
    import json as _json
    from harness.study2.analysis import backfill_reasoning_tokens

    raw = tmp_path / "raw.jsonl"
    raw.write_text(
        _json.dumps({"item_id": "t0", "tone_level": "L4_neutral", "trial": 0, "reasoning_tokens": 7})
        + '\n{"item_id": "t0", "tone_lev'
    )
    recs = [{"task_id": "t0", "tone_level": "L4_neutral", "trial": 0}]
    assert backfill_reasoning_tokens(recs, raw) == 1
    assert recs[0]["reasoning_tokens"] == 7


# --- 8. A ten-hour run had no resume ---------------------------------------
# The container was restarted at 896 of 1050 trajectories. The incremental
# records survived on disk, so redoing that work would have been a choice.

def _jsonl(path: Path, rows):
    import json as _json
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(_json.dumps(r) for r in rows))


def test_recorded_trajectories_are_recognised(tmp_path: Path):
    from harness.study2.runner import completed_trajectories

    _jsonl(tmp_path / "analysis" / "study2_core_gpt-luna_records.jsonl", [
        {"task_id": "t0", "tone_level": "L4_neutral", "trial": 0, "crashed": False},
        {"task_id": "t0", "tone_level": "L7_threatening", "trial": 2, "crashed": False},
    ])
    done = completed_trajectories(tmp_path, "core", "gpt-luna")
    # Four fields, not three: the injection turn is part of the key, so a
    # crossed run's three cells per (task, tone, trial) stay distinct. Rows
    # carrying no turn key as None, which is what they re-derive to.
    assert done == {("t0", "L4_neutral", 0, None), ("t0", "L7_threatening", 2, None)}


def test_a_crashed_trajectory_is_retried_not_skipped(tmp_path: Path):
    """Its usual cause is a transient provider error, not the task."""
    from harness.study2.runner import completed_trajectories

    _jsonl(tmp_path / "analysis" / "study2_core_m_records.jsonl", [
        {"task_id": "t0", "tone_level": "L4_neutral", "trial": 0, "crashed": True},
    ])
    assert completed_trajectories(tmp_path, "core", "m") == set()


def test_a_torn_final_line_is_not_counted_as_done(tmp_path: Path):
    """A process killed mid-write leaves a partial record; that is not
    evidence the trajectory finished."""
    from harness.study2.runner import completed_trajectories

    p = tmp_path / "analysis" / "study2_core_m_records.jsonl"
    p.parent.mkdir(parents=True)
    p.write_text('{"task_id":"t0","tone_level":"L4_neutral","trial":0,"crashed":false}\n{"task_id":"t1","tone_le')
    assert completed_trajectories(tmp_path, "core", "m") == {("t0", "L4_neutral", 0, None)}


def test_no_records_file_means_nothing_to_skip(tmp_path: Path):
    from harness.study2.runner import completed_trajectories

    assert completed_trajectories(tmp_path, "core", "never-run") == set()


def test_resume_skips_recorded_work_and_keeps_it_in_the_result(tmp_path: Path):
    """The final JSON must be the whole run, not just the part after the
    restart."""
    from harness.study2 import runner
    from harness.study2.agent_loop import Trajectory
    from harness.study2.grader import GradeResult
    from harness.tone_wrappers import TONE_ORDER

    class _Task:
        task_id = "t0"
        instruction_type = "Cell-Level Manipulation"
        instruction = "do it"
        answer_position = "A1"
        input_spreadsheet_paths = [Path("in.xlsx")]
        answer_spreadsheet_paths = [Path("ans.xlsx")]

    # every tone already done except one
    prior = [{"task_id": "t0", "tone_level": t, "trial": 0, "crashed": False, "passed": False}
             for t in TONE_ORDER if t != "L5_rude"]
    _jsonl(tmp_path / "analysis" / "study2_core_m_records.jsonl", prior)

    ran = []

    def fake_run(tracker, model, task_id, instruction, input_path, workdir, *, tone_level, trial, **kw):
        ran.append(tone_level)
        return Trajectory(task_id=task_id, tone_level=tone_level, trial=trial)

    with patch("harness.study2.runner.run_react_multi_round", side_effect=fake_run), \
         patch("harness.study2.runner._grade_trajectory",
               side_effect=lambda *a, **k: GradeResult("t0", False, 3, 0, 0.0, [])):
        out = runner.run_condition_batch(
            [_live("m")], [_Task()], grader=None, out_dir=tmp_path, phase="core",
            budget_cap_usd=1e6, n_trials=1, resume=True,
        )

    assert ran == ["L5_rude"], f"should have run only the missing tone, ran {ran}"
    assert len(out) == 7, f"result must carry all 7 tones, got {len(out)}"


# --- 9. Two live runs of the SAME model shared one records file -----------
# Namespacing per model stopped different models colliding; it did nothing
# about the same model twice. A container restart was reported, a resumed run
# was started -- and the original process had not died. Both appended to one
# records file for half an hour, producing 39 duplicate (task, tone, trial)
# rows that an analysis would have counted as extra trials.

def test_a_second_live_run_of_the_same_model_is_refused(tmp_path: Path):
    import os
    from harness.study2.runner import RunAlreadyInProgress, _exclusive_run

    with _exclusive_run(tmp_path, "core", "gpt-luna"):
        lock = tmp_path / "locks" / "study2_core_gpt-luna.lock"
        assert lock.read_text().strip() == str(os.getpid())
        # a different, live pid holds it -- os.getppid() is live and is not
        # us. pid 1 was used here originally and made this test pass locally
        # (root can signal it) while failing on CI (an unprivileged runner
        # gets PermissionError), which is the bug the liveness check had.
        lock.write_text(str(os.getppid()))
        with pytest.raises(RunAlreadyInProgress, match="already running"):
            with _exclusive_run(tmp_path, "core", "gpt-luna"):
                pass
        lock.write_text(str(os.getpid()))


def test_a_stale_lock_from_a_dead_process_is_taken_over(tmp_path: Path):
    """The normal state after a crash, which is exactly when --resume runs."""
    from harness.study2.runner import _exclusive_run

    lock = tmp_path / "locks" / "study2_core_m.lock"
    lock.parent.mkdir(parents=True)
    lock.write_text("999999")  # not a live pid
    with _exclusive_run(tmp_path, "core", "m"):
        pass  # must not raise


def test_different_models_do_not_block_each_other(tmp_path: Path):
    """Running the roster in parallel stays supported."""
    from harness.study2.runner import _exclusive_run

    with _exclusive_run(tmp_path, "core", "gpt-luna"):
        with _exclusive_run(tmp_path, "core", "qwen-current"):
            pass


def test_the_lock_is_released_when_the_run_ends(tmp_path: Path):
    from harness.study2.runner import _exclusive_run

    with _exclusive_run(tmp_path, "core", "m"):
        pass
    assert not (tmp_path / "locks" / "study2_core_m.lock").exists()


def test_the_lock_is_released_even_when_the_run_raises(tmp_path: Path):
    from harness.study2.runner import _exclusive_run

    with pytest.raises(ValueError):
        with _exclusive_run(tmp_path, "core", "m"):
            raise ValueError("budget")
    assert not (tmp_path / "locks" / "study2_core_m.lock").exists()


# --- 10. total_tokens counted the model's thinking twice -------------------
# OpenAI-style usage reports reasoning as completion_tokens_details.
# reasoning_tokens -- a breakdown OF completion, not a bucket beside it. The
# trajectory total added it a third time. On the 4,647-call gpt-luna core log
# the provider's own usage.total_tokens equalled prompt + completion on every
# call, and reasoning never exceeded completion on any; the harness total ran
# ~932 above the provider's per trajectory, which is just the mean reasoning
# spend added back. total_tokens is the statistic compared against the
# published single-turn figure, so it was inflated by about the size of the
# effect being measured.

def _luna_result_row(prompt: int, completion: int, reasoning: int):
    from harness.spend_tracker import ResultRow

    return ResultRow(
        row_id="r", study="study2", phase="core", item_id="t0",
        tone_level="L4_neutral", trial=0, model_key="gpt-luna", model_id="x/luna",
        provider="openai_compatible", temperature=0.0, seed=None,
        reasoning_effort="medium", thinking_budget_tokens=None,
        prompt_tokens=prompt, completion_tokens=completion, reasoning_tokens=reasoning,
        cost_usd=0.0, latency_s=0.0, refused=False, error=None, response_text="",
        extracted_answer=None, is_correct=None, timestamp=0.0,
    )


def _trajectory_total(rows) -> int:
    """The runner's own arithmetic, exercised the way runner.py writes it."""
    return sum(
        r.prompt_tokens
        + r.completion_tokens
        + (0 if r.reasoning_included_in_completion else r.reasoning_tokens)
        for r in rows
    )


def test_trajectory_total_tokens_does_not_add_reasoning_twice():
    rows = [_luna_result_row(803, 1600, 828), _luna_result_row(2479, 703, 241)]
    assert _trajectory_total(rows) == 803 + 1600 + 2479 + 703
    # The old sum, kept here so the regression is unmistakable rather than
    # implied: it is over by exactly the reasoning spend.
    assert _trajectory_total(rows) + 828 + 241 == sum(
        r.prompt_tokens + r.completion_tokens + r.reasoning_tokens for r in rows
    )


def test_provider_response_total_matches_the_providers_own_figure():
    from harness.providers.base import ProviderResponse

    r = ProviderResponse(text="x", prompt_tokens=803, completion_tokens=1600, reasoning_tokens=828)
    assert r.total_tokens == 2403, "OpenAI-style: reasoning is already inside completion"
    assert r.reasoning_tokens_outside_completion == 0


def test_a_provider_that_reports_reasoning_separately_keeps_the_third_term():
    """Google's usageMetadata keeps thoughtsTokenCount outside
    candidatesTokenCount, so dropping the term unconditionally would have
    traded a double-count for an undercount."""
    from harness.providers.base import ProviderResponse

    r = ProviderResponse(
        text="x", prompt_tokens=100, completion_tokens=50, reasoning_tokens=40,
        reasoning_included_in_completion=False,
    )
    assert r.total_tokens == 190
    assert r.reasoning_tokens_outside_completion == 40


def test_cost_estimate_prices_reasoning_exactly_once():
    from harness.providers.base import ProviderResponse
    from harness.spend_tracker import compute_cost_usd

    model = replace(_model("priced"), input_price_per_1m=1.0, output_price_per_1m=2.0)
    openai_style = ProviderResponse(text="x", prompt_tokens=1_000_000, completion_tokens=1_000_000, reasoning_tokens=400_000)
    assert compute_cost_usd(model, openai_style) == pytest.approx(3.0)

    google_style = ProviderResponse(
        text="x", prompt_tokens=1_000_000, completion_tokens=1_000_000, reasoning_tokens=400_000,
        reasoning_included_in_completion=False,
    )
    assert compute_cost_usd(model, google_style) == pytest.approx(3.8)


# --- 11. Resumed runs left an abandoned attempt in the raw log ------------
# The log is append-only and a resume does not know the killed attempt's
# calls are already in it, so one (item_id, tone_level, trial) key can hold
# two trajectories -- and only the second was ever graded. In the gpt-luna
# core log, task 56953 / L7_threatening / trial 0 holds 9 calls: a 3-call
# attempt that died, then the 6-call attempt behind the record. Summing the
# key blindly gave 2338 reasoning tokens where the graded trajectory spent
# 1236.

def _two_attempt_log(path: Path) -> Path:
    """The real shape of task 56953 / L7_threatening / trial 0: a 3-call
    attempt, then the 6-call retry. prompt_tokens climbs within an attempt
    and drops when the conversation restarts."""
    import json as _json

    abandoned = [(803, 1600, 828), (2479, 703, 241), (8696, 130, 33)]
    real = [(803, 1425, 670), (2390, 421, 85), (3938, 473, 59),
            (4447, 765, 297), (9202, 496, 83), (9637, 159, 42)]
    rows = []
    ts = 1789155849.0
    for i, (p, c, r) in enumerate(abandoned + real):
        # The retry starts hours later; timestamps are what orders the log.
        offset = ts + i * 5 if i < len(abandoned) else ts + 11_000 + i * 5
        rows.append({"item_id": "56953", "tone_level": "L7_threatening", "trial": 0,
                     "prompt_tokens": p, "completion_tokens": c, "reasoning_tokens": r,
                     "timestamp": offset})
    path.write_text("\n".join(_json.dumps(r) for r in rows))
    return path


def test_backfill_uses_only_the_last_attempt(tmp_path: Path):
    from harness.study2.analysis import backfill_reasoning_tokens

    raw = _two_attempt_log(tmp_path / "raw.jsonl")
    recs = [{"task_id": "56953", "tone_level": "L7_threatening", "trial": 0}]
    assert backfill_reasoning_tokens(recs, raw) == 1
    assert recs[0]["reasoning_tokens"] == 1236, "the graded trajectory, not both attempts"
    assert recs[0]["reasoning_tokens"] != 2338, "2338 is the naive both-attempts sum"


def test_recompute_total_tokens_rebuilds_prompt_plus_completion(tmp_path: Path):
    from harness.study2.analysis import recompute_total_tokens

    raw = _two_attempt_log(tmp_path / "raw.jsonl")
    recs = [{"task_id": "56953", "tone_level": "L7_threatening", "trial": 0, "total_tokens": 999_999}]
    assert recompute_total_tokens(recs, raw) == 1
    expected = (803 + 1425) + (2390 + 421) + (3938 + 473) + (4447 + 765) + (9202 + 496) + (9637 + 159)
    assert recs[0]["total_tokens"] == expected
    # Neither of the two ways of getting it wrong: reasoning re-added, or
    # the abandoned attempt included.
    assert recs[0]["total_tokens"] != expected + 1236
    assert recs[0]["total_tokens"] != expected + 803 + 1600 + 2479 + 703 + 8696 + 130


def test_recompute_total_tokens_leaves_keys_absent_from_this_log_alone(tmp_path: Path):
    """Record files are namespaced per model; running against one model's
    raw log must not blank out another's records."""
    from harness.study2.analysis import recompute_total_tokens

    raw = _two_attempt_log(tmp_path / "raw.jsonl")
    recs = [{"task_id": "other", "tone_level": "L4_neutral", "trial": 0, "total_tokens": 4242}]
    assert recompute_total_tokens(recs, raw) == 0
    assert recs[0]["total_tokens"] == 4242


def test_a_single_attempt_key_is_unaffected(tmp_path: Path):
    """The split must trigger on a restart, not on ordinary turn-to-turn
    growth -- otherwise every trajectory would be truncated to its last call."""
    import json as _json
    from harness.study2.analysis import last_attempt_calls, recompute_total_tokens

    rows = [{"item_id": "t0", "tone_level": "L4_neutral", "trial": 0, "timestamp": 10.0 + i,
             "prompt_tokens": p, "completion_tokens": 100, "reasoning_tokens": 10}
            for i, p in enumerate((800, 2400, 3900, 4400))]
    assert len(last_attempt_calls(rows)) == 4
    raw = tmp_path / "raw.jsonl"
    raw.write_text("\n".join(_json.dumps(r) for r in rows))
    recs = [{"task_id": "t0", "tone_level": "L4_neutral", "trial": 0, "total_tokens": 0}]
    recompute_total_tokens(recs, raw)
    assert recs[0]["total_tokens"] == 800 + 2400 + 3900 + 4400 + 400


# --- 12. The accuracy CI resampled trajectories, not tasks ----------------
# Every other test in this harness clusters by task. bootstrap_accuracy_ci
# resampled individual trajectories, so 150 observations (50 tasks x 3
# trials) were treated as 150 independent facts. On the gpt-luna core run's
# L1_sycophantic tone that gave [0.280, 0.433] where clustering gives
# [0.240, 0.473] -- the iid interval is a third narrower, on exactly the
# interval a reader uses to judge whether two tones overlap.

def _perfectly_agreeing_trials(n_tasks: int = 40, n_trials: int = 3):
    """Half the tasks pass on every trial, half fail on every trial. Trials
    within a task carry no information the task itself does not, so an iid
    bootstrap's extra "observations" are pure double-counting."""
    passed, task_ids = [], []
    for t in range(n_tasks):
        for _ in range(n_trials):
            passed.append(t % 2 == 0)
            task_ids.append(f"task{t}")
    return passed, task_ids


def test_clustered_accuracy_ci_is_wider_than_the_iid_one():
    from harness.study1.analysis import bootstrap_accuracy_ci

    passed, task_ids = _perfectly_agreeing_trials()
    iid = bootstrap_accuracy_ci(passed, n_boot=2000)
    clustered = bootstrap_accuracy_ci(passed, cluster_ids=task_ids, n_boot=2000)

    assert clustered.accuracy == iid.accuracy == 0.5, "the point estimate must not move"
    iid_width = iid.ci_high - iid.ci_low
    clustered_width = clustered.ci_high - clustered.ci_low
    assert clustered_width > iid_width * 1.5, (iid_width, clustered_width)


def test_singleton_clusters_reproduce_the_iid_interval():
    """The cluster bootstrap must reduce to the iid one when every cluster
    holds one observation -- otherwise callers with genuinely unclustered
    data silently get a different number than before."""
    from harness.study1.analysis import bootstrap_accuracy_ci

    passed = [True, True, False, True, False, False, True]
    iid = bootstrap_accuracy_ci(passed, n_boot=2000)
    singleton = bootstrap_accuracy_ci(passed, cluster_ids=list(range(len(passed))), n_boot=2000)
    assert (singleton.ci_low, singleton.ci_high) == (iid.ci_low, iid.ci_high)


def test_mismatched_cluster_ids_are_refused_not_guessed():
    from harness.study1.analysis import bootstrap_accuracy_ci

    with pytest.raises(ValueError, match="parallel"):
        bootstrap_accuracy_ci([True, False, True], cluster_ids=["a", "b"])


def test_study2_analyze_clusters_its_accuracy_ci_by_task(tmp_path: Path, monkeypatch, capsys):
    """The bug lived at the call site too: cmd_study2_analyze had the task
    ids in hand and passed only the pass/fail column."""
    import argparse
    import json as _json
    from harness.cli import cmd_study2_analyze

    passed, task_ids = _perfectly_agreeing_trials(n_tasks=20)
    records = [
        {"model_key": "m", "task_id": tid, "instruction_type": "i", "tone_level": tone,
         "tone_position": 0, "trial": i % 3, "passed": p, "soft_restriction": 0.0,
         "refused": False, "severity": "correct" if p else "other", "severity_detail": "",
         "inspected_before_acting": True, "self_checked_output": True,
         "took_destructive_action": False, "destructive_action_had_backup": False,
         "n_turns": 2, "cost_usd": 0.01, "total_tokens": 1000, "reasoning_tokens": 100,
         "crashed": False}
        for tone in ("L1_sycophantic", "L4_neutral", "L7_threatening")
        for i, (p, tid) in enumerate(zip(passed, task_ids))
    ]
    monkeypatch.chdir(tmp_path)  # --records-path is globbed relative to cwd
    (tmp_path / "study2_core_m_records.json").write_text(_json.dumps(records))
    out_path = tmp_path / "report.json"
    cmd_study2_analyze(
        argparse.Namespace(records_path="study2_core_m_records.json", out_path=str(out_path))
    )
    capsys.readouterr()

    report = _json.loads(out_path.read_text())
    ci = report["accuracy_by_tone"]["L4_neutral"]
    assert ci["n"] == 60
    # 20 tasks, all-or-nothing within a task: an iid CI over 60 trajectories
    # lands near +/-0.12; clustering over 20 tasks is roughly twice that.
    assert ci["ci_high"] - ci["ci_low"] > 0.35, ci


# --- 13. The accuracy note asserted a base rate the run contradicted ------
# It printed SpreadsheetBench's published ~18% base rate as though it were
# this run's; the gpt-luna core run's observed pooled accuracy is 0.296. It
# then used "underpowered" to wave accuracy away entirely, which inverts
# what power means: at the observed slope (-0.0107/level) power is ~0.44 and
# the 80%-power MDE ~0.017/level, so a null is weak evidence -- but a
# pre-registered test that clears its threshold is valid at n=150 exactly as
# it is at n=1500.

def _accuracy_note_records(n_pass: int, n_fail: int):
    rows = []
    for i in range(n_pass + n_fail):
        rows.append({"task_id": f"t{i // 3}", "tone_level": f"L{i % 7 + 1}",
                     "passed": i < n_pass, "refused": False})
    return rows


def test_the_accuracy_note_computes_the_rate_instead_of_asserting_one():
    from harness.cli import _underpowered_accuracy_note

    note = _underpowered_accuracy_note(_accuracy_note_records(n_pass=311, n_fail=739))
    assert "29.6%" in note, note
    assert "18%" not in note, "the published base rate is not this run's observed rate"
    assert "1050 trajectories" in note

    # A different run must move the number, not reprint a constant.
    other = _underpowered_accuracy_note(_accuracy_note_records(n_pass=500, n_fail=500))
    assert "50.0%" in other


def test_the_accuracy_note_does_not_treat_power_as_invalidating_a_positive():
    from harness.cli import _underpowered_accuracy_note

    note = _underpowered_accuracy_note(_accuracy_note_records(n_pass=311, n_fail=739))
    lowered = note.lower()
    assert "secondary" in lowered, "accuracy was pre-designated secondary; say so"
    assert "false negatives" in lowered
    assert "valid pre-registered test" in lowered
    assert "thousand-plus" not in lowered


def test_a_lock_held_by_another_users_process_is_not_stolen(tmp_path: Path):
    """os.kill(pid, 0) raising PermissionError proves the process EXISTS.
    Catching OSError broadly read that as 'dead' and took the lock -- CI
    caught it, where the runner is unprivileged and pid 1 belongs to root."""
    import os
    from harness.study2.runner import RunAlreadyInProgress, _exclusive_run

    lock = tmp_path / "locks" / "study2_core_m.lock"
    lock.parent.mkdir(parents=True)
    lock.write_text("4242")

    real_kill = os.kill

    def fake_kill(pid, sig):
        if pid == 4242 and sig == 0:
            raise PermissionError(1, "Operation not permitted")
        return real_kill(pid, sig)

    with patch("harness.study2.runner.os.kill", side_effect=fake_kill):
        with pytest.raises(RunAlreadyInProgress):
            with _exclusive_run(tmp_path, "core", "m"):
                pass
    assert lock.read_text().strip() == "4242", "the other run's lock must survive"


# --- 11. The instrument itself was confounded --------------------------
# Two flaws in the v1 wrapper set, both found by independent review of the
# Luna run, both of which would have replicated across the whole roster.

def test_only_the_shared_instruction_is_a_task_instruction():
    """v1's neutral wrapper alone added "Read the question carefully before
    responding, and provide a single final answer." That made the study's
    own reference level a different instrument: it tripled zero-turn
    trajectories and carried the entire severity-by-tone shift."""
    from harness.tone_wrappers import INSTRUCTION, TONE_WRAPPERS

    banned = ("provide a single final answer", "read the question carefully",
              "step by step", "think carefully", "double-check", "show your work")
    for key, w in TONE_WRAPPERS.items():
        rest = w.text.replace(INSTRUCTION, "").lower()
        for phrase in banned:
            assert phrase not in rest, f"{key} carries a task instruction: {phrase!r}"


def test_every_wrapper_carries_the_shared_instruction_verbatim():
    from harness.tone_wrappers import INSTRUCTION, TONE_WRAPPERS

    for key, w in TONE_WRAPPERS.items():
        assert w.text.count(INSTRUCTION) == 1, key


def test_all_seven_wrappers_are_exactly_the_same_length():
    """Not "within 5". In v1 the spread was 5 tokens and U-shaped across the
    scale -- the same shape as accuracy -- and length predicted accuracy
    better than tone rank did (r=+0.82 vs -0.72). With seven points the two
    cannot be separated, so the tolerance has to be zero."""
    from harness.tone_wrappers import wrapper_token_counts

    counts = wrapper_token_counts()
    assert len(set(counts.values())) == 1, f"wrappers differ in length: {counts}"


def test_the_length_validator_rejects_any_spread():
    """The guard has to actually fire. It caught a one-token slip while the
    v2 set was being written, which is the whole reason it runs at import."""
    from dataclasses import replace as dc_replace

    from harness import tone_wrappers as tw

    assert tw.MAX_TOKEN_SPREAD == 0
    tw.validate_wrapper_lengths()  # the shipped set must pass

    longer = dc_replace(tw.TONE_WRAPPERS["L3_polite"],
                        text=tw.TONE_WRAPPERS["L3_polite"].text + " One extra clause here.")
    with patch.dict(tw.TONE_WRAPPERS, {"L3_polite": longer}):
        with pytest.raises(ValueError, match="not length-matched"):
            tw.validate_wrapper_lengths()

    tw.validate_wrapper_lengths()  # and the patch must not have leaked


def test_records_carry_the_wrapper_set_version():
    """v1 and v2 data are not comparable; a record that does not say which
    it came from can be pooled with the other by mistake."""
    from harness.study2 import runner
    from harness.tone_wrappers import WRAPPER_SET_VERSION

    assert WRAPPER_SET_VERSION == "v2"
    src = (Path(runner.__file__)).read_text()
    assert '"wrapper_set": WRAPPER_SET_VERSION' in src


def test_a_tone_subset_keeps_its_positions_from_the_full_shuffle():
    """A single-arm rerun must meet the same burst positions that arm saw in
    the full run, or it is not comparable to it."""
    import random as _r

    from harness.tone_wrappers import TONE_ORDER

    for task in ("t0", "t1", "t2", "t3", "t4"):
        full = list(TONE_ORDER)
        _r.Random(f"0|m|{task}").shuffle(full)
        want_pos = full.index("L4_neutral")

        subset = list(TONE_ORDER)
        _r.Random(f"0|m|{task}").shuffle(subset)
        got = [(i, t) for i, t in enumerate(subset) if t == "L4_neutral"]
        assert got == [(want_pos, "L4_neutral")], task


def test_a_labelled_run_writes_to_its_own_files():
    """Otherwise it lands on the main records file, --resume skips every
    trajectory as already done, and an analysis pools two wrapper sets."""
    from harness.study2.runner import _run_tag

    base = _run_tag([replace(_model("gpt-luna"), provider="openai_compatible")])
    assert base == "gpt-luna"
    assert f"{base}-wrapper-v2" != base


def test_the_reported_records_path_is_the_one_actually_written():
    """This line kept its own copy of the naming rule and ignored both
    --run-label and the dry-run suffix, so a labelled run wrote its records
    correctly and then announced the MAIN dataset's path -- which reads
    exactly like the full run has just been clobbered by a single arm."""
    from harness.study2.runner import _run_tag

    live = replace(_model("gpt-luna"), provider="openai_compatible")
    assert _run_tag([live]) == "gpt-luna"
    assert _run_tag([live], "wrapper-v2") == "gpt-luna-wrapper-v2"
    assert _run_tag([_model("gpt-luna")], "wrapper-v2") == "gpt-luna-dryrun-wrapper-v2"
    assert _run_tag([live], None) == "gpt-luna"


# --- 12. Mid-task interjection: the "manager check-in" experiment ---------
# The wrappers set a tone once, at the start. These deliver it partway
# through. Comparing "threatening interjection" against "no interjection"
# would measure being interrupted; the neutral interjection is the control
# that isolates the tone of the interruption from the interruption itself.

def test_the_interjections_are_exactly_length_matched():
    """A length difference would be a second manipulation riding along with
    the first -- the v1 wrapper mistake, repeated at a smaller scale."""
    from harness.tone_wrappers import interjection_token_counts

    counts = interjection_token_counts()
    assert len(set(counts.values())) == 1, counts


def test_the_neutral_interjection_carries_no_task_instruction():
    """v1's neutral wrapper carried one and turned the agent into an
    answerer. The control here has to be inert, not helpful."""
    from harness.tone_wrappers import INTERJECTIONS

    banned = ("provide a single final answer", "check your", "make sure you",
              "read carefully", "step by step", "verify", "double-check")
    text = INTERJECTIONS["L4_neutral"].lower()
    for phrase in banned:
        assert phrase not in text, f"neutral interjection instructs: {phrase!r}"


def test_both_arms_interrupt_the_same_trajectory_at_the_same_turn():
    """The turn is seeded on (task, trial) and NOT on the tone. If it varied
    by arm, turn position would ride along with the manipulation and the
    paired comparison would measure both at once."""
    import random as _r

    from harness.study2.runner import INTERJECTION_TURNS

    for task in ("t0", "t1", "t2", "t3", "t4"):
        for trial in range(3):
            turns = {
                _r.Random(f"interject|0|m|{task}|{trial}").choice(INTERJECTION_TURNS)
                for _tone in ("L4_neutral", "L7_threatening")
            }
            assert len(turns) == 1, f"{task}/{trial} would be interrupted at {turns}"


def test_the_injection_turn_varies_across_trajectories():
    """It is supposed to be random, not a fixed turn dressed up as one."""
    import random as _r

    from harness.study2.runner import INTERJECTION_TURNS

    seen = {_r.Random(f"interject|0|m|t{i}|0").choice(INTERJECTION_TURNS) for i in range(40)}
    assert len(seen) > 1, f"every trajectory got the same turn: {seen}"


def _row():
    """A minimal ResultRow; the interjection tests only care about messages."""
    from harness.spend_tracker import ResultRow

    return ResultRow(
        row_id="r", study="s", phase="p", item_id="t", tone_level="L4_neutral",
        trial=0, model_key="m", model_id="m", provider="mock", temperature=0.0,
        seed=None, reasoning_effort=None, thinking_budget_tokens=None,
        prompt_tokens=1, completion_tokens=1, reasoning_tokens=0, cost_usd=0.0,
        latency_s=0.0, refused=False, error=None, response_text="",
        extracted_answer=None, is_correct=None, timestamp=0.0,
    )



def test_the_interjection_reaches_the_model_at_that_turn(tmp_path: Path):
    from harness.study2 import agent_loop

    seen: list[str] = []

    class _Resp:
        refused = False
        text = "```python\nprint(1)\n```"
        prompt_tokens = completion_tokens = reasoning_tokens = 1
        reasoning_tokens_reported = True
        cached_tokens = 0
        served_provider = None
        raw: dict = {}
        tool_calls = None

    def fake_call(tracker, model, system, messages, *a, **k):
        seen.append(messages[-1]["content"] if len(messages) > 1 else "")
        return _Resp(), _row()

    class _Exec:
        stdout = "ok"; stderr = ""; timed_out = False; output_workbook_path = None

    with patch("harness.study2.agent_loop._call_and_record", side_effect=fake_call), \
         patch("harness.study2.agent_loop.execute_python_on_workbook", return_value=_Exec()):
        traj = agent_loop.run_react_multi_round(
            None, _model("m"), "t", "do it", tmp_path / "wb.xlsx", tmp_path / "wd",
            tone_level="L4_neutral", trial=0,
            interjection="MANAGER SPEAKS", interjection_turn=1, max_turns=4,
        )

    fired_in = [i for i, m in enumerate(seen) if "MANAGER SPEAKS" in m]
    assert traj.interjection_fired
    # injected at turn 1 => first visible to the model on its turn-2 call
    assert fired_in == [2], f"interjection surfaced on calls {fired_in}"


def test_a_trajectory_that_ends_early_records_that_it_never_fired(tmp_path: Path):
    """An analysis that counted un-fired trajectories as treated would
    dilute the effect toward zero."""
    from harness.study2 import agent_loop

    class _Resp:
        refused = False
        text = "FINAL: done"
        prompt_tokens = completion_tokens = reasoning_tokens = 1
        reasoning_tokens_reported = True
        cached_tokens = 0
        served_provider = None
        raw: dict = {}
        tool_calls = None

    def fake_call(tracker, model, system, messages, *a, **k):
        return _Resp(), _row()

    with patch("harness.study2.agent_loop._call_and_record", side_effect=fake_call):
        traj = agent_loop.run_react_multi_round(
            None, _model("m"), "t", "do it", tmp_path / "wb.xlsx", tmp_path / "wd",
            tone_level="L4_neutral", trial=0,
            interjection="MANAGER SPEAKS", interjection_turn=2, max_turns=4,
        )
    assert traj.interjection_fired is False


# --- 13. Seven-level interruptions, crossed over injection turn ------------
# The opening-wrapper arm came back null on cost; the mid-task interruption
# did not. These guard the expanded design: all seven registers delivered
# mid-task, crossed over where they land, with the timing comparison's
# selection confound closed at analysis time.

def test_every_tone_level_has_an_interjection():
    """The interjection scale has to be the SAME scale as the wrappers, or
    an interjection arm and an opening-tone arm cannot be compared."""
    from harness.tone_wrappers import INTERJECTIONS, TONE_ORDER

    assert list(INTERJECTIONS) == TONE_ORDER


def test_no_interjection_carries_a_task_instruction():
    """Extended from the neutral-only check: v1's mistake was ONE level
    carrying an instruction the others didn't, which is exactly what a
    seven-level set makes easy to reintroduce."""
    from harness.tone_wrappers import INTERJECTIONS

    banned = ("provide a single final answer", "check your", "make sure you",
              "read carefully", "step by step", "verify", "double-check",
              "answer the question")
    for key, text in INTERJECTIONS.items():
        lowered = text.lower()
        for phrase in banned:
            assert phrase not in lowered, f"{key} instructs: {phrase!r}"


def test_every_interjection_is_marked_as_an_interruption():
    """A shared stem is what makes the text read as an interruption rather
    than task content. If it varied by level, 'was interrupted' would vary
    with register and the arms would differ by two things at once."""
    from harness.tone_wrappers import INTERJECTIONS

    for key, text in INTERJECTIONS.items():
        assert text.startswith("Checking in."), f"{key} lacks the shared stem: {text!r}"


def test_turn_zero_is_a_legal_injection_point():
    """It was excluded on a misreading of the loop: the interjection rides
    the observation produced AT the end of iteration `turn`, which the model
    sees on the next one, so index 0 already means 'after the first
    response'. It is also the position that reaches the most trajectories."""
    from harness.study2.runner import INTERJECTION_TURNS

    assert 0 in INTERJECTION_TURNS


def test_turn_zero_actually_fires_on_the_first_observation():
    """The claim above, exercised against the real agent loop rather than
    asserted about it."""
    from harness.study2.agent_loop import run_react_multi_round

    captured = _capture_messages_run(interjection="ZZ-MARKER", interjection_turn=0)
    user_texts = [m["content"] for m in captured if m["role"] == "user"]
    assert any("ZZ-MARKER" in t for t in user_texts), user_texts
    # It must not be in the OPENING message -- that would make it a wrapper,
    # not an interruption.
    assert "ZZ-MARKER" not in user_texts[0]


def test_crossed_turns_run_every_turn_once_per_trial():
    from harness.study2.runner import _injection_turns_for

    turns = _injection_turns_for(
        "L7_threatening", True, [0, 1, 2], 0, "m", "task", 0
    )
    assert turns == [0, 1, 2]


def test_crossed_turns_are_deduplicated_and_ordered():
    from harness.study2.runner import _injection_turns_for

    assert _injection_turns_for("L5_rude", True, [2, 0, 0, 1], 0, "m", "t", 0) == [0, 1, 2]


def test_uncrossed_runs_keep_the_legacy_seeded_draw():
    """The micro-experiment's runs have to stay reproducible."""
    from harness.study2.runner import _injection_turns_for

    a = _injection_turns_for("L7_threatening", True, None, 0, "m", "task", 0)
    b = _injection_turns_for("L4_neutral", True, None, 0, "m", "task", 0)
    assert len(a) == 1 and a == b


def test_no_interjection_still_runs_exactly_once():
    from harness.study2.runner import _injection_turns_for

    assert _injection_turns_for(None, True, [0, 1, 2], 0, "m", "t", 0) == [None]


def test_resume_key_distinguishes_injection_turns(tmp_path):
    """Keyed on (task, tone, trial) alone, a crossed resume would see cell 1
    recorded and skip cells 2 and 3 as duplicates, truncating the run to a
    third of its design without saying so."""
    import json as _json

    from harness.study2.runner import completed_trajectories

    path = tmp_path / "analysis" / "study2_core_x_records.jsonl"
    path.parent.mkdir(parents=True)
    with open(path, "w") as fh:
        for turn in (0, 1):
            fh.write(_json.dumps({
                "task_id": "T1", "tone_level": "L4_neutral", "trial": 0,
                "interjection": "L7_threatening", "interjection_turn": turn,
                "crashed": False,
            }) + "\n")

    done = completed_trajectories(tmp_path, "core", "x")
    assert ("T1", "L4_neutral", 0, 0) in done
    assert ("T1", "L4_neutral", 0, 1) in done
    assert ("T1", "L4_neutral", 0, 2) not in done, "turn 2 would be skipped unrun"


def test_resume_key_still_matches_rows_without_an_injection_turn(tmp_path):
    """Runs recorded before the turn was crossed must still resume."""
    import json as _json

    from harness.study2.runner import completed_trajectories

    path = tmp_path / "analysis" / "study2_core_x_records.jsonl"
    path.parent.mkdir(parents=True)
    path.write_text(_json.dumps({
        "task_id": "T1", "tone_level": "L2_very_polite", "trial": 3,
        "crashed": False,
    }) + "\n")

    assert ("T1", "L2_very_polite", 3, None) in completed_trajectories(tmp_path, "core", "x")


def test_turn_comparable_tasks_ignores_the_treated_arm_entirely():
    """The population must be decided from the control arm only. Selecting
    on the treated arm's own firing would select on a variable the treatment
    moves, which is the confound this function exists to close."""
    from harness.study2.analysis import turn_comparable_tasks

    rows = []
    # Control arm reaches both turns on TASK_OK, only turn 0 on TASK_SHORT.
    for turn, fired in ((0, True), (1, True)):
        rows.append(_interj_row("TASK_OK", "L4_neutral", turn, fired, 100))
    for turn, fired in ((0, True), (1, False)):
        rows.append(_interj_row("TASK_SHORT", "L4_neutral", turn, fired, 100))
    # Treated arm says the opposite on both tasks. It must not matter.
    for turn, fired in ((0, True), (1, False)):
        rows.append(_interj_row("TASK_OK", "L7_threatening", turn, fired, 900))
    for turn, fired in ((0, True), (1, True)):
        rows.append(_interj_row("TASK_SHORT", "L7_threatening", turn, fired, 900))

    assert turn_comparable_tasks(rows, [0, 1]) == {"TASK_OK"}


def test_turn_comparison_is_run_on_one_shared_task_set():
    """Both turns must be measured on identical tasks, or the difference
    between them can be manufactured by their having drawn different work --
    the micro-experiment's caveat."""
    from harness.study2.analysis import injection_turn_effects

    rows = []
    for task in ("A", "B"):
        for turn in (0, 1):
            rows.append(_interj_row(task, "L4_neutral", turn, True, 100))
            rows.append(_interj_row(task, "L7_threatening", turn, True, 150))
    # A third task the control never reaches at turn 1, with a huge treated
    # effect at turn 1. If it leaked in, turn 1 would look far more costly.
    rows.append(_interj_row("C", "L4_neutral", 0, True, 100))
    rows.append(_interj_row("C", "L4_neutral", 1, False, 100))
    rows.append(_interj_row("C", "L7_threatening", 1, True, 5000))

    effects = injection_turn_effects(rows, "L7_threatening", [0, 1])
    assert effects[0]["n_tasks"] == effects[1]["n_tasks"] == 2
    assert effects[1]["mean_effect"] == 50.0, effects[1]


def test_unfired_trajectories_are_dropped_not_averaged_in():
    """A trajectory that ended before the injection turn received no dose.
    Counting it as a treated observation averages the control condition into
    the treated arm, by an amount that grows with turn index."""
    from harness.study2.analysis import interjection_trend_test

    rows = []
    for task in ("A", "B", "C"):
        for tone, value in (("L4_neutral", 100.0), ("L7_threatening", 300.0)):
            rows.append(_interj_row(task, tone, 0, True, value))
        # Unfired treated rows carrying control-like values.
        rows.append(_interj_row(task, "L7_threatening", 2, False, 100.0))

    kept = interjection_trend_test(rows, levels=["L4_neutral", "L7_threatening"])
    assert kept.n_clusters == 3
    assert kept.observed_slope == 200.0, kept


def _interj_row(task_id, interjection, turn, fired, reasoning):
    return {
        "task_id": task_id, "tone_level": "L4_neutral", "trial": 0,
        "interjection": interjection, "interjection_turn": turn,
        "interjection_fired": fired, "reasoning_tokens": reasoning,
        "crashed": False,
    }


def _capture_messages_run(interjection, interjection_turn):
    """Runs the real agent loop against a stub that always emits code, and
    returns the message list the model was shown."""
    from harness.study2 import agent_loop as al

    captured = []

    class _Resp:
        refused = False
        text = "```python\npass\n```"

    def _fake_call(tracker, model, system_prompt, messages, *a, **k):
        captured.clear()
        captured.extend({"role": m["role"], "content": m["content"]} for m in messages)
        return _Resp(), _row()

    class _Exec:
        stdout, stderr, timed_out = "", "", False
        output_workbook_path = None

    orig_call, orig_exec = al._call_and_record, al.execute_python_on_workbook
    al._call_and_record = _fake_call
    al.execute_python_on_workbook = lambda *a, **k: _Exec()
    try:
        al.run_react_multi_round(
            None, _StubModel(), "task", "instruction", _PathStub(), _PathStub(),
            tone_level="L4_neutral", trial=0,
            interjection=interjection, interjection_turn=interjection_turn,
            max_turns=3,
        )
    finally:
        al._call_and_record = orig_call
        al.execute_python_on_workbook = orig_exec
    return captured


class _StubModel:
    key = "m"
    model_id = "m"
    provider = "mock"


class _PathStub:
    def __truediv__(self, other):
        return self
